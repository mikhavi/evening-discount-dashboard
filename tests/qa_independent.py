"""Independent read-only audit of the supplied source and sample output.

Run with a Python environment that has openpyxl installed. This intentionally
does not import or reuse the browser application's calculation code.
"""

from collections import Counter
from datetime import datetime, time
from decimal import Decimal, ROUND_CEILING, ROUND_FLOOR
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "01_выгрузка_остатков_и_продаж_1С.xlsx"
OUTPUT = ROOT / "Итоги_Вечерней_Уценки_20_00.xlsx"
ALLOWED = {
    "Кулинария: Салаты заправленные",
    "Кулинария: Салаты незаправленные",
    "Кулинария: Горячие блюда",
    "Кулинария: Выпечка и сдоба",
}
D = lambda value: Decimal(str(value))


def hours(value):
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return D(value.hour) + D(value.minute) / 60 + D(value.second) / 3600
    hour, minute, *second = str(value).split(":")
    return D(hour) + D(minute) / 60 + D(second[0] if second else 0) / 3600


def floor5(value):
    return (value / 5).to_integral_value(rounding=ROUND_FLOOR) * 5


def ceil5(value):
    return (value / 5).to_integral_value(rounding=ROUND_CEILING) * 5


source = load_workbook(SOURCE, read_only=True, data_only=True)
rows = list(source.active.values)
assert len(rows) - 1 == 107

counts = Counter()
old_total = Decimal(0)
promo_total = Decimal(0)
cost_total = Decimal(0)
before_loss = Decimal(0)
during_loss = Decimal(0)
promo_rows = []
for row in rows[1:]:
    code, barcode, name, category, stock, unit, retail, cost, made, shelf = row[:10]
    category = " ".join(category.strip().replace(":", ": ").split()) if category else ""
    if category not in ALLOWED:
        counts["out_of_scope"] += 1
        continue
    counts["target"] += 1
    stock, retail, cost = D(stock), D(retail), D(cost)
    expiry = hours(made) + D(shelf)
    if unit == "шт" and stock != stock.to_integral_value():
        counts["fractional_pieces"] += 1
        continue
    low = stock <= (D("0.5") if unit == "кг" else D(1))
    if low:
        counts["low_stock"] += 1
    if expiry <= 20:
        counts["before"] += 1
        before_loss += stock * cost
    elif expiry <= 23:
        counts["during"] += 1
        during_loss += stock * cost
    if low and expiry <= 20:
        counts["low_and_before"] += 1
    if low or expiry <= 23:
        continue
    price30 = floor5(retail * D("0.7"))
    minimum = cost * D("1.05")
    new_price = price30 if price30 >= minimum else ceil5(minimum)
    if new_price >= retail:
        counts["blocked_margin"] += 1
        continue
    if new_price != price30:
        counts["margin_alert"] += 1
    if not barcode or "E+" in str(barcode).upper() or ";" in str(barcode):
        counts["barcode_review_in_plan"] += 1
    counts["promo"] += 1
    old_total += stock * retail
    promo_total += stock * new_price
    cost_total += stock * cost
    promo_rows.append((name, str(barcode or "").strip(), category, stock, new_price, stock * new_price))

assert counts == Counter({
    "out_of_scope": 31,
    "target": 76,
    "promo": 54,
    "before": 19,
    "during": 2,
    "low_stock": 1,
    "low_and_before": 1,
    "fractional_pieces": 1,
    "margin_alert": 6,
    "barcode_review_in_plan": 7,
})
assert (old_total, promo_total, cost_total) == (D(729005), D("536765.5"), D(429816))
assert (before_loss, during_loss) == (D("156502.5"), D(19990))

output = load_workbook(OUTPUT, read_only=False, data_only=True)
assert output.sheetnames == [
    "Сводка", "План 19-40", "Цена скорректирована", "Бракераж",
    "Исключения", "Проверка целевых", "Срок годности",
]
plan = output["План 19-40"]
assert plan.max_row - 1 == 54
assert len(output["Цена скорректирована"]["A"]) - 1 == 6
assert len(output["Бракераж"]["A"]) - 1 == 1
assert plan.freeze_panes == "A2"
assert plan.auto_filter.ref == "A1:H55"
assert all(plan.cell(r, 2).data_type == "s" for r in range(2, 56))
assert all(plan.cell(r, 2).number_format == "@" for r in range(2, 56))
assert abs(sum(D(plan.cell(r, 5).value) * D(plan.cell(r, 6).value) for r in range(2, 56)) - promo_total) < D("0.000001")
barcode_mismatches = []
for r in range(2, 56):
    name, barcode, category, old, new, stock, unit, mark = [plan.cell(r, c).value for c in range(1, 9)]
    expected = next((item for item in promo_rows if item[0] == name), None)
    assert expected is not None, (r, name)
    if (barcode or "") != expected[1]:
        barcode_mismatches.append((r, name, barcode, expected[1]))
    assert (category, D(stock), D(new)) == (expected[2], expected[3], expected[4]), (r, name)
    assert D(new) * D(stock) == expected[5], (r, name)

print("BARCODE_MISMATCHES", [(r, ascii(name), ascii(actual), ascii(expected)) for r, name, actual, expected in barcode_mismatches])
assert not barcode_mismatches
print("PASS: independent source calculations and sample workbook agree")
print(dict(counts))
print({"old_value": str(old_total), "potential": str(promo_total), "cost": str(cost_total), "before_loss": str(before_loss), "during_loss": str(during_loss)})
